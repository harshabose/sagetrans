"""`tables.xlsx`: every table of the results folder, as data (docs/OUTPUTS_spec.md 6).

Built from the loaded tables, nothing recomputed. Sheet `index` first; then one sheet per CSV with the
CSV header in row 1 (the column's identity), the unit in row 2, data from row 3. Numbers are numbers,
booleans are TRUE/FALSE, an empty cell is empty. No formulas, no charts, no macros. The document
properties are fixed (creator `sagetrans`, created and modified = the run date) so two builds from
the same folder have the same sheets.
"""

from __future__ import annotations

import datetime as dt
import math
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.writer.excel import ExcelWriter

from sagetrans.io.view.load import Results, Table
from sagetrans.io.view.units import excel_format, unit_of

DESCRIPTIONS = {
    "A_feasibility": "A: zero-loss feasibility over entry speed and deceleration",
    "A_max_feasible_decel": "A: largest and smallest feasible deceleration by scenario and speed",
    "A_window": "A: the deceleration window at each entry speed and what closes it",
    "Q2_pusher_margin": "Q2: pusher acceleration by altitude, mass and state of charge",
    "Q2_worst": "Q2: the worst pusher case",
    "D_recommendation": "D: the recommended Q_TRANS_DECEL and its ensemble statistics",
    "D_recommendation_ensemble": "D: one row per ensemble member",
    "D_a_rec_curve": "D: achievable deceleration against ground speed",
    "D_trade_curve": "D: distance and energy against planned deceleration",
    "F_descent": "F: descent rate cap and its binding limit",
    "F_profile": "F: descent profile",
    "E_summary": "E: range, reserves and flyability by site and mission",
    "E_budget": "E: energy by segment for each range",
    "E_segments": "E: every segment of every mission",
    "E_scenarios": "E: what each scenario is and its AIRBRAKE-exit state at the register and recommended AIRSPEED_MIN",
    "C_params": "C: the optimum, the register start and the ArduPilot default",
    "C_runs": "C: one row per ensemble member and case",
    "C_search": "C: best cost of each search seed",
    "C_q_curve": "C: the optimum's ramp with Q_TRANS_DECEL varied",
    "C_sens": "C: one-at-a-time sensitivity at the optimum",
    "C_history": "C: time series of member 0 in each case",
    "C_alpha_curve": "C: the price of the margin, one search per wing limit",
    "C_ref_params": "C unconstrained (reference only): the optimum, the register start, the default",
    "C_ref_runs": "C unconstrained (reference only): one row per ensemble member and case",
    "C_ref_history": "C unconstrained (reference only): the optimum on the history scenario",
    "FT_params": "FT: the optimum, the register start and the ArduPilot default",
    "FT_runs": "FT: one row per member and case",
    "FT_search": "FT: best cost of each search seed",
    "FT_scan": "FT: best cost with AIRSPEED_MIN fixed on a grid",
    "FT_sens": "FT: one-at-a-time sensitivity at the optimum",
    "FT_robust": "FT: AIRSPEED_MIN against the C_L,max scale (reporting only)",
    "FT_history": "FT: time series of the named scenarios",
    "sens_oat": "Sensitivity: one-at-a-time swings",
    "sens_global": "Sensitivity: Morris screening",
    "settings_recommended": "The recommended ArduPilot settings sheet",
    "inputs_register": "Every register entry as used, with its evidence grade",
}


def _cell(v: Any) -> Any:
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _sheet_names(tables: list[Table]) -> dict[str, str]:
    """File stem -> sheet name: truncated to 31 characters and unique (the index has the full name)."""
    used: set[str] = set()
    out: dict[str, str] = {}
    for t in tables:
        base = t.name[:31]
        name, k = base, 1
        while name.lower() in used or name.lower() == "index":
            suffix = f"~{k}"
            name = base[: 31 - len(suffix)] + suffix
            k += 1
        used.add(name.lower())
        out[t.name] = name
    return out


def _run_date(res: Results) -> dt.datetime:
    try:
        return dt.datetime.fromisoformat(str(res.record.get("date")))
    except ValueError:
        return dt.datetime(2000, 1, 1)


def write_workbook(res: Results, path: Path) -> None:
    """Write `tables.xlsx` for a loaded results folder."""
    tables = [t for n, t in res.tables.items() if n not in res.aliased]
    names = _sheet_names(tables)
    wb = Workbook()
    when = _run_date(res)
    wb.properties.creator = "sagetrans"
    wb.properties.lastModifiedBy = "sagetrans"
    wb.properties.created = when
    wb.properties.modified = when
    wb.properties.title = f"sagetrans results {res.folder.name}"
    ix = wb.active
    assert ix is not None
    ix.title = "index"
    r = res.record
    draft = res.label == "DRAFT"
    ix["A1"] = "DRAFT — not for release" if draft else f"sagetrans results {res.folder.name}"
    ix["A1"].font = Font(bold=True, color="FFC00000" if draft else "FF000000", size=14)
    for i, (k, v) in enumerate(
        (
            ("run label", res.label),
            ("config hash", r.get("config_hash", "")),
            ("commit", f"{r.get('git_commit', '')}{'+dirty' if r.get('git_dirty') else ''}"),
            ("date", r.get("date", "")),
        ),
        start=2,
    ):
        ix.cell(row=i, column=1, value=k).font = Font(bold=True)
        ix.cell(row=i, column=2, value=str(v))
    head_row = 7
    for j, h in enumerate(("sheet", "source csv", "rows", "columns", "analysis", "description"), 1):
        ix.cell(row=head_row, column=j, value=h).font = Font(bold=True)
    for i, t in enumerate(tables, start=head_row + 1):
        sheet = names[t.name]
        c = ix.cell(row=i, column=1, value=sheet)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1")
        c.font = Font(color="FF0563C1", underline="single")
        ix.cell(row=i, column=2, value=t.source)
        ix.cell(row=i, column=3, value=len(t.rows))
        ix.cell(row=i, column=4, value=len(t.columns))
        ix.cell(row=i, column=5, value=t.name.split("_")[0])
        ix.cell(row=i, column=6, value=DESCRIPTIONS.get(t.name, ""))
    for col, w in zip("ABCDEF", (30, 34, 8, 9, 10, 70), strict=True):
        ix.column_dimensions[col].width = w
    for t in tables:
        ws = wb.create_sheet(names[t.name])
        for j, col in enumerate(t.columns, 1):
            ws.cell(row=1, column=j, value=col).font = Font(bold=True)
            ws.cell(row=2, column=j, value=unit_of(col)[0] or None)
        fmts = [excel_format(c) for c in t.columns]
        for i, row in enumerate(t.rows, start=3):
            for j, col in enumerate(t.columns, 1):
                v = _cell(row.get(col))
                cell = ws.cell(row=i, column=j, value=v)
                if isinstance(v, float) and fmts[j - 1] != "General":
                    cell.number_format = fmts[j - 1]
        for j, col in enumerate(t.columns, 1):
            width = max([len(col)] + [len(str(rw.get(col, ""))) for rw in t.rows[:200]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), 40)
        ws.freeze_panes = "A3"
        if t.columns:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(t.columns))}{max(len(t.rows) + 2, 2)}"
    path.parent.mkdir(parents=True, exist_ok=True)
    # `Workbook.save` stamps the modified time with the clock; the writer underneath does not, so
    # two builds from one folder have the same properties (spec 6.8)
    ExcelWriter(wb, ZipFile(path, "w", ZIP_DEFLATED, allowZip64=True)).save()
