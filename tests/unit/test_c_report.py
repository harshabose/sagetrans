"""The C chapter (and the FT chapter) of the results viewer: scenarios, margins, spreads."""

import csv
import json
import math
import re
from pathlib import Path

import numpy as np

from sagetrans.io import stages
from sagetrans.io.view import charts as ch
from sagetrans.io.view.chapters import (
    history_panels,
    scenario_caption,
    select_history,
)
from sagetrans.io.view.load import load
from sagetrans.io.view.registry import build_html

ROOT = Path(__file__).resolve().parents[2]


def write_table(folder: Path, name: str, rows: list[dict]) -> None:
    cols: list[str] = []
    for r in rows:
        cols += [c for c in r if c not in cols]
    with (folder / f"{name}.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def write_record(folder: Path, **extra) -> None:
    rec = {"label": "DRAFT", "reasons": ["72 register entries are TBD"], "config_hash": "ab" * 32,
           "git_commit": "deadbee", "git_dirty": False, "date": "2026-10-29", "seeds": [1, 2],
           "outputs": {}, "analyses": {}, "search_spreads": {}, "schema_version": 2, **extra}
    (folder / "run_record.json").write_text(json.dumps(rec), encoding="utf-8")


def section(html: str, chapter: str) -> str:
    return html.split(f'id="{chapter}"')[1].split("</section>")[0]


# ---- a synthetic run: C on two scenarios and two members --------------------------------------
def hist_rows(case, view, scenario, member, v0=15.0, n=6, **extra):
    meta = {"view": view, "scenario": scenario, "altitude_m": 0.0, "isa_offset_k": 0.0, "soc": 0.8,
            "headwind_ms": 0.0, "mass_kg": 26.7, "V0_tas_ms": v0, "V0_eas_ms": v0 * 0.99,
            "band_eps_h_m": 1.0, "run_overshoot_m": 1.5 if "default" not in case else 3.1}
    return [{"case": case, "member": member, **meta, **extra, "t": 0.5 * i, "x": 10.0 * i,
             "h": 0.1 * i, "V": v0 - 2.0 * i, "v_h": 0.0, "theta": 4.0 * i, "alpha_w_deg": 3.0 * i,
             "delta_r": 0.05 * i, "theta_cmd": 4.0 * i, "omega_r": 100.0}
            for i in range(n)]


def synthetic_run(folder: Path) -> None:
    write_table(folder, "C_params", [
        {"case": "optimum", "T_bt_s": 1.6, "theta_A_deg": 19.6, "spin_min": 0.11,
         "Q_TRANS_DECEL_ms2": 2.03, "J": 1.48, "p95_distance_m": 119.0, "mean_energy_kJ": 20.0},
        {"case": "ardupilot_default", "T_bt_s": 3.0, "theta_A_deg": 30.0, "spin_min": 0.15,
         "Q_TRANS_DECEL_ms2": 2.0, "J": 61.5, "p95_distance_m": 121.0, "mean_energy_kJ": 20.3}])
    c_hist = (hist_rows("C optimum", "core", "S-01", -1)
              + hist_rows("C ardupilot_default", "core", "S-01", -1, v0=15.0)
              + hist_rows("C optimum", "core", "S-03", -1, v0=19.0)  # another scenario
              + hist_rows("C optimum", "default_fails", "ensemble member", 14, v0=19.3,
                          default_fails=True)
              + hist_rows("C ardupilot_default", "default_fails", "ensemble member", 14, v0=19.3,
                          default_fails=True)
              + hist_rows("C optimum", "default_fails", "ensemble member", 2, v0=17.0,
                          default_fails=True))  # another member
    write_table(folder, "C_history", c_hist)
    write_record(folder)


def numeric(rows):
    return [{k: (float(v) if re.fullmatch(r"-?[0-9.]+", v or "x") else v) for k, v in r.items()}
            for r in rows]


# ---- C1 ---------------------------------------------------------------------------------------
def test_a_history_figure_never_mixes_scenarios_or_members(tmp_path):
    synthetic_run(tmp_path)
    c_rows = numeric(csv.DictReader((tmp_path / "C_history.csv").open(encoding="utf-8")))
    core, ident = select_history(c_rows, "core", "S-01")
    assert ident == ("scenario", "S-01")
    assert set(core) == {"C optimum", "C ardupilot_default"}
    assert {r["scenario"] for rows in core.values() for r in rows} == {"S-01"}  # not S-03
    worst, ident = select_history(c_rows, "default_fails")
    assert ident == ("member", 14.0) and set(worst) == {"C optimum", "C ardupilot_default"}
    assert {r["member"] for rows in worst.values() for r in rows} == {14.0}  # not member 2


def test_the_c_chapter_has_two_history_figures_with_a_stated_scenario(tmp_path):
    synthetic_run(tmp_path)
    c = section(build_html(load(tmp_path)), "c")
    assert "C on one scenario (S-01)" in c
    assert "Where the default fails (ensemble member 14)" in c
    assert "Member 0 of the ensemble" not in c  # the old single figure is gone
    # the caption states altitude, ISA, mass, SoC, headwind and the entry speed, true and equivalent
    for text in ("altitude 0 m", "ISA +0.0 K", "mass 26.7 kg", "SoC 0.80", "headwind +0.00 m/s",
                 "entry airspeed 15.0 m/s true, 14.8 m/s equivalent"):
        assert text in c, text
    assert "C optimum (15.0 m/s)" in c  # the entry speed in the legend
    assert "B1" not in c
    assert "stops at" in c  # an end marker on every curve
    assert "+1 m" in c and "-1 m" in c  # the altitude band as two lines
    assert "overshoots this member by 3.10 m" in c  # the default's overshoot on that member


def test_the_airspeed_axis_starts_at_zero_and_a_run_ends_with_a_marker():
    t = list(range(6))
    s = ch.TSeries("a", "s1", "a", [float(x) for x in t],
                   {"V": [15.0 - 2.0 * x for x in t], "h": [0.0] * 6, "theta": [0.0] * 6,
                    "delta_r": [0.0] * 6}, end_mark=True)
    svg = ch.panel_figure("f", history_panels(1.0), [s])
    v_panel = svg.split('data-pi="0"')[1].split("</svg>")[0]
    ticks = re.findall(r'class="tick"[^>]*>([^<]*)<', v_panel)
    assert "0" in ticks and all(not tk.startswith("-") for tk in ticks)  # not -10
    assert svg.count("stops at") == 4  # one marker per panel


def test_scenario_caption_skips_what_is_missing():
    assert scenario_caption({"altitude_m": 500.0}) == "altitude 500 m"
    assert "entry airspeed" not in scenario_caption({"V0_tas_ms": 15.0})


def test_the_history_rows_end_at_the_stop():
    assert stages._sample_indices(12, 5) == [0, 5, 10, 11]
    assert stages._sample_indices(11, 5) == [0, 5, 10]
    assert stages._sample_indices(1, 5) == [0]

    class Sim:
        t = np.array([0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.6])

        def __getitem__(self, k):
            return np.zeros(len(self.t))

    rows = stages._c_trace(Sim(), 0.0, "C optimum", -1, meta={"scenario": "S-01"})
    assert [r["t"] for r in rows] == [0.0, 0.5, 0.6] and rows[-1]["scenario"] == "S-01"
    assert math.isfinite(rows[0]["V"])


def test_a_default_with_no_failing_member_says_so(tmp_path):
    synthetic_run(tmp_path)
    rows = list(csv.DictReader((tmp_path / "C_history.csv").open(encoding="utf-8")))
    for r in rows:
        if r["view"] == "default_fails":
            r["default_fails"] = "False"
    write_table(tmp_path, "C_history", rows)
    c = section(build_html(load(tmp_path)), "c")
    assert "No member fails under the default; this is the one with the largest overshoot" in c


# ---- C2 ---------------------------------------------------------------------------------------
class CfgStub:
    os_tol_m, sink_tol_ms = 2.0, 1.76
    margin_flag_m, margin_flag_ms, margin_frac = 0.5, 0.3, 0.75
    q_fine_offsets_ms2 = [0.0]


def test_margins_to_both_tolerances_and_the_edge_flag():
    m = stages.margin_cols(CfgStub, 1.59, 1.24)
    assert m["os_margin_m"] == 2.0 - 1.59 and m["sink_margin_ms"] == 1.76 - 1.24
    assert m["at_edge"] is True  # the overshoot margin 0.41 m is below 0.5 m
    assert stages.margin_cols(CfgStub, 1.0, 1.0)["at_edge"] is False
    assert stages.margin_cols(CfgStub, 0.5, 1.6)["at_edge"] is True  # the SINK margin alone: 0.16
    assert stages.margin_cols(CfgStub, 0.5, 1.6)["sink_margin_ms"] < 0.3
    empty = stages.margin_cols(CfgStub, "", "")  # nothing stopped: no margin, no edge
    assert empty["os_margin_m"] == "" and empty["sink_margin_ms"] == ""
    assert empty["at_edge"] is False
    assert empty["os_tol_m"] == 2.0 and empty["sink_tol_ms"] == 1.76


def fine(q, os_, sink):
    return {"q": q, "max_overshoot_m": os_, "max_end_sink_ms": sink, "J": 1.5 + q}


def test_the_margin_row_is_the_largest_q_within_both_fractions():
    rows = [fine(1.93, 1.0, 1.0), fine(2.03, 1.3, 1.2), fine(2.08, 1.45, 1.3),
            fine(2.13, 1.6, 1.0), fine(2.23, 3.0, 1.0)]  # limits: 1.5 m and 1.32 m/s
    assert stages.pick_margin_row(rows, CfgStub)["q"] == 2.08
    # a q that meets the overshoot margin but not the sink margin is not the margin row
    rows2 = [fine(2.03, 1.0, 1.0), fine(2.08, 1.2, 1.5)]
    assert stages.pick_margin_row(rows2, CfgStub)["q"] == 2.03


def test_no_q_on_the_grid_meets_the_margin():
    rows = [fine(2.03, 1.8, 1.0), fine(2.08, 1.9, 1.0), fine(2.13, 2.5, 1.0)]
    assert stages.pick_margin_row(rows, CfgStub) is None
    assert stages.pick_margin_row([], CfgStub) is None
    assert stages.pick_margin_row([fine(2.0, "", "")], CfgStub) is None  # nothing stopped


def test_the_new_c_settings_are_refused_when_out_of_range():
    import pytest
    from pydantic import ValidationError

    from sagetrans.io.settings import CCfg

    assert CCfg().margin_frac == 0.75 and CCfg().q_fine_offsets_ms2[2] == 0.0  # the defaults
    for bad in ({"margin_frac": 0.0}, {"margin_frac": 1.5}, {"margin_flag_m": -0.1},
                {"margin_flag_ms": -1.0}, {"q_fine_offsets_ms2": [0.0, float("nan")]},
                {"q_fine_offsets_ms2": [9.0]}):
        with pytest.raises(ValidationError):
            CCfg(**bad)


def test_the_q_note_is_computed_from_the_numbers():
    from types import SimpleNamespace

    from sagetrans.io.recommended import q_note

    inp = SimpleNamespace(study=SimpleNamespace(C=CfgStub))
    c_row = {"Q_TRANS_DECEL_ms2": 2.03, "max_overshoot_m": 1.59, "max_end_sink_ms": 1.24}
    d_row = {"a_rec_min_ms2": 1.65}
    m_row = {"Q_TRANS_DECEL_ms2": 1.98}
    text = q_note(inp, c_row, d_row, m_row)
    assert "C's q is 2.03 m/s^2" in text and "D's a_rec minimum is 1.65 m/s^2" in text
    assert "C is 23 % above it" in text  # 100 * (2.03 / 1.65 - 1)
    assert "worst overshoot is 1.59 m against the tolerance 2 m (PROPOSED)" in text
    assert "worst end sink is 1.24 m/s against 1.76 m/s (PROPOSED)" in text
    assert "margin row" in text and "q = 1.98 m/s^2" in text and "75%" in text
    # D not in the run, and no q on the grid meets the margin: both stated, nothing invented
    text2 = q_note(inp, c_row, None, None)
    assert "D is not in this run" in text2 and "no q on the fine grid keeps both" in text2


def c2_run(folder: Path, edge: bool = True, with_margin: bool = True) -> None:
    opt = {"case": "optimum", "T_bt_s": 1.6, "theta_A_deg": 19.6, "spin_min": 0.11,
           "Q_TRANS_DECEL_ms2": 2.03, "J": 1.48, "p95_distance_m": 119.0, "mean_energy_kJ": 20.0,
           "max_overshoot_m": 1.59, "max_end_sink_ms": 1.24, "os_tol_m": 2.0, "sink_tol_ms": 1.76,
           "os_margin_m": 0.41, "sink_margin_ms": 0.52, "at_edge": edge}
    params = [opt, {**opt, "case": "ardupilot_default", "J": 61.5, "Q_TRANS_DECEL_ms2": 2.0,
                    "max_overshoot_m": 3.8, "os_margin_m": -1.8, "at_edge": True}]
    if with_margin:
        params.append({**opt, "case": "optimum_margin", "Q_TRANS_DECEL_ms2": 1.98,
                       "max_overshoot_m": 1.4, "os_margin_m": 0.6, "at_edge": False})
    write_table(folder, "C_params", params)
    curve = [{"case": f"q = {q:g}", "q": q, "grid": "coarse", "J": 1.5, "n_overshoot": 0.0,
              "n_sinking": 0.0, "max_overshoot_m": 0.4 * q, "max_end_sink_ms": 0.6}
             for q in (1.0, 1.5, 2.0, 2.5)]
    curve += [{"case": f"q = {q:.3f}", "q": q, "grid": "fine", "J": 1.5,
               "n_overshoot": 0.0 if q < 2.1 else 11.0, "n_sinking": 0.0,
               "max_overshoot_m": 1.59 + 5 * (q - 2.03), "max_end_sink_ms": 1.2}
              for q in (1.98, 2.03, 2.08, 2.13, 2.18)]
    write_table(folder, "C_q_curve", curve)
    write_record(folder)


def test_the_c_chapter_shows_both_margins_the_chart_and_the_proposed_caveat(tmp_path):
    c2_run(tmp_path)
    c = section(build_html(load(tmp_path)), "c")
    assert "Worst overshoot / tolerance" in c and "1.59 / 2" in c
    assert "Worst end sink / tolerance" in c and "1.24 / 1.76" in c
    assert "the optimum is at the edge of a tolerance: overshoot margin 0.41 m" in c
    assert "Worst overshoot and worst end sink against q" in c
    assert c.count("tolerance 2 (PROPOSED)") == 1 and "tolerance 1.76 (PROPOSED)" in c  # both lines
    assert "relies on the overshoot and end-sink tolerances, which are PROPOSED" in c
    assert "margin row (reported, not recommended)" in c


def test_the_margin_row_is_never_the_recommendation(tmp_path):
    from sagetrans.io.inputs import load_inputs
    from sagetrans.io.recommended import settings_rows

    inp = load_inputs(ROOT / "tests" / "data" / "placeholder_configs", only={"C"})
    base = {"T_bt_s": 1.6, "theta_A_deg": 19.6, "spin_min": 0.11, "at_bound": ""}
    tables = {"C_params": [
        {"case": "optimum", "Q_TRANS_DECEL_ms2": 2.03, **base, "max_overshoot_m": 1.59,
         "max_end_sink_ms": 1.24},
        {"case": "optimum_margin", "Q_TRANS_DECEL_ms2": 1.5, **base},
        {"case": "ardupilot_default", "Q_TRANS_DECEL_ms2": 2.0, **base}]}
    rows = settings_rows(inp, tables)
    q = next(r for r in rows if r["parameter"] == "Q_TRANS_DECEL" and r["analysis"] == "C")
    assert float(q["value"]) == 2.03  # the optimum, not the margin row
    assert "margin row" in q["note"] and "q = 1.50 m/s^2" in q["note"]  # but it is named
    assert "D is not in this run" in q["note"]  # the tiny run has no D: said, not broken


# ---- C3 ---------------------------------------------------------------------------------------
def test_a_strip_shows_the_second_case_as_hollow_dots_and_a_tolerance_rule():
    dots = [(0.8, "m0", False, "s1"), (1.2, "m1", False, "s1")]
    plain = ch.strip_chart([("overshoot [m]", dots, None)], "x")
    full = ch.strip_chart([("overshoot [m]", dots, None)], "x",
                          extras=[([(3.4, "d0"), (3.8, "d1")], (2.0, "tolerance 2"))])
    assert "mk ref hollow" not in plain
    assert full.count("mk ref hollow") == 2 and "tolerance 2" in full
    def top(svg):  # the right-hand end label of the strip's axis
        return max(float(v) for v in re.findall(r'class="tick"[^>]*>([-0-9.]+)<', svg))

    assert top(plain) < 1.5 and top(full) >= 3.8  # the axis grew to include the default's members


def test_the_spread_statement_is_computed_and_names_the_cause():
    from sagetrans.io.view.chapters import spread_statement

    rows = [{"overshoot_m": v, "end_sink_ms": 0.5 + 0.1 * i}
            for i, v in enumerate((0.77, 1.1, 1.73))]
    s = spread_statement(rows, {})
    assert "No member of the optimum stops short" in s and "positive in all 3" in s
    assert "0.77 to 1.73 m" in s and "at most 0.70 m/s" in s
    assert "planned deceleration q from t = 0" in s and "pitch ramp over T_bt" in s
    short = spread_statement([{"overshoot_m": -0.2}, {"overshoot_m": 0.5}], {})
    assert "1 of 2 members stop short" in short
    assert spread_statement([], {}) == ""


def test_the_c_chapter_strips_carry_the_default_members_and_the_tolerances(tmp_path):
    c2_run(tmp_path)
    runs = []
    for case, base in (("optimum", 1.0), ("ardupilot_default", 3.0)):
        for m in range(4):
            runs.append({"case": case, "member": m, "energy_kJ": 20.0, "excursion_m": 1.0,
                         "end_sink_ms": 0.5, "overshoot_m": base + 0.2 * m, "flags": ""})
    write_table(tmp_path, "C_runs", runs)
    c = section(build_html(load(tmp_path)), "c")
    assert c.count("mk ref hollow") == 16  # four default members on each of the four strips
    assert "tolerance 2" in c and "tolerance 1.76" in c
    assert "member of the ArduPilot default" in c
    assert "No member of the optimum stops short" in c


# ---- C4 ---------------------------------------------------------------------------------------
class FakeRun:
    def __init__(self, failures, stopped=True, vh_end=0.1, overshoot=0.5):
        self.metrics = type("M", (), {"failures": tuple(failures)})()
        self.stopped, self.vh_end, self.overshoot = stopped, vh_end, overshoot


def test_failure_counts_are_complete_and_the_smooth_cost_drops_the_penalty():
    runs = [FakeRun([]), FakeRun(["overshoot"]), FakeRun(["sinking"]),
            FakeRun(["sinking", "overshoot"]), FakeRun(["not_stopped"], stopped=False)]
    cnt = stages.failure_counts(runs)
    assert cnt == {"n_failing_members": 4.0, "n_not_stopped": 1.0, "n_fail_total": 5.0}
    assert stages.failure_counts([FakeRun([])]) == {  # a clean case still has every key, zeros
        "n_failing_members": 0.0, "n_not_stopped": 0.0, "n_fail_total": 0.0}
    w = type("W", (), {"penalty": 10.0})()
    assert stages.smooth_cost(61.5, 5.0, w) == 61.5 - 50.0
    assert stages.smooth_cost(1.48, 0.0, w) == 1.48


def sens_rows():
    base = {"n_overshoot": 0.0, "n_sinking": 0.0, "n_other_flags": 0.0, "n_failing_members": 0.0}
    return [
        {"parameter": "rotor_lag_s", "value": 0.25, **base, "J": 1.48, "J_smooth": 1.48},
        {"parameter": "rotor_lag_s", "value": 0.5, **base, "J": 111.0, "J_smooth": 1.6,
         "n_sinking": 11.0, "n_failing_members": 11.0},
        {"parameter": "z_acc_p", "value": 0.015, **base, "J": 61.0, "J_smooth": 1.5,
         "n_sinking": 6.0, "n_failing_members": 6.0},
    ]


def test_the_sensitivity_shows_the_failing_members_apart_from_the_smooth_change(tmp_path):
    c2_run(tmp_path)
    rows = sens_rows()
    write_table(tmp_path, "C_sens", rows)
    params = list(csv.DictReader((tmp_path / "C_params.csv").open(encoding="utf-8")))
    for r in params:
        r["J_smooth"] = "1.48"
    write_table(tmp_path, "C_params", params)
    c = section(build_html(load(tmp_path)), "c")
    assert "Members that fail" in c and "Change of J without the failure penalty" in c
    assert "x rotor_lag_s = 0.5" in c and "x z_acc_p = 0.015" in c  # the failing rows are marked
    assert "x rotor_lag_s = 0.25" not in c
    assert "11 failing member(s): sinking 11" in c and "6 failing member(s): sinking 6" in c
    assert "(1.600 / 1.480 - 1) x 100" in c  # the expression stays in the tooltip
    assert "+7500" not in c  # the penalty no longer sets the scale
    # an older results folder (no J_smooth) still gets the single chart
    for r in rows:
        r.pop("J_smooth")
    write_table(tmp_path, "C_sens", rows)
    for r in params:
        r.pop("J_smooth")
    write_table(tmp_path, "C_params", params)
    old = section(build_html(load(tmp_path)), "c")
    assert "Members that fail" not in old and "One-at-a-time sensitivity" in old


# ---- C5 ---------------------------------------------------------------------------------------
def energy_figure(c_html: str) -> str:
    return c_html.split("Energy against stopping distance")[1].split("</svg>")[0]


def test_the_energy_figure_labels_every_q_and_draws_the_optimum_last(tmp_path):
    c2_run(tmp_path)
    curve = list(csv.DictReader((tmp_path / "C_q_curve.csv").open(encoding="utf-8")))
    for r, (d, e) in zip(curve[:4], ((80, 10), (95, 14), (119, 20), (150, 26)), strict=False):
        r["p95_distance_m"], r["mean_energy_kJ"] = d, e
    write_table(tmp_path, "C_q_curve", curve)
    params = list(csv.DictReader((tmp_path / "C_params.csv").open(encoding="utf-8")))
    for r in params:  # the optimum and the default almost on top of each other
        near = (119.5, 20.1) if r["case"] == "optimum" else (121, 20.3)
        r["p95_distance_m"], r["mean_energy_kJ"] = near
    write_table(tmp_path, "C_params", params)
    fig = energy_figure(section(build_html(load(tmp_path)), "c"))
    for q in ("q = 1", "q = 1.5", "q = 2", "q = 2.5"):  # every point of the coarse curve
        assert f">{q}<" in fig, q
    assert "q = 1.98" not in fig and "fine" not in fig  # the fine grid has its own chart
    assert fig.index("C optimum") > fig.index("ArduPilot default")  # the optimum is drawn last
    # the two labels near each other are nudged apart, not stacked
    found = re.findall(r'class="dl" x="[0-9.]+" y="([0-9.]+)"[^>]*>([^<]*)<', fig)
    ys = {text: float(y) for y, text in found}
    assert abs(ys["C optimum"] - ys["ArduPilot default"]) >= 20.0


# ---- C6 ---------------------------------------------------------------------------------------
def seeds(js, **cols):
    return [{"seed": i + 1, "best_cost": j, "evaluations": 160,
             **{k: v[i] for k, v in cols.items()}} for i, j in enumerate(js)]


def test_the_seed_statement_tells_a_ridge_from_a_single_optimum_from_no_convergence():
    from sagetrans.io.view.chapters import seed_statement

    ridge = seeds([1.4796, 1.4789, 1.4795], T_bt_s=[1.2, 1.6, 2.1],
                  theta_A_deg=[19.5, 19.6, 19.6], spin_min=[0.11, 0.11, 0.11],
                  Q_TRANS_DECEL_ms2=[2.03, 2.03, 2.04])
    s = seed_statement(ridge)
    assert "agree on J" in s and "flat ridge" in s and "T_bt 1.2 to 2.1" in s
    assert "theta_A" not in s and "spin_min" not in s  # only what actually differs is named
    same = seeds([1.4796, 1.4789], T_bt_s=[1.6, 1.61], theta_A_deg=[19.6, 19.6],
                 spin_min=[0.109, 0.11], Q_TRANS_DECEL_ms2=[2.03, 2.03])
    assert "single well-defined optimum" in seed_statement(same)
    loose = seeds([1.40, 1.55], T_bt_s=[1.6, 1.6])
    assert "disagree on J" in seed_statement(loose) and "not converged" in seed_statement(loose)
    assert seed_statement(seeds([1.4796, 1.4789])) == ""  # an older folder: no parameters
    assert seed_statement(seeds([1.4])) == ""


def test_the_bound_chip_says_where_the_upper_bound_comes_from():
    from sagetrans.io.view.chapters import bound_reason

    row = {"theta_A_cfg_hi_deg": 30.0, "ptch_lim_deg": 20.0, "theta_A_hi_deg": 20.0}
    text = bound_reason(row, "theta_A:hi")
    assert "the upper bound of theta_A is the pitch limit" in text
    assert "min(30, PTCH_LIM_MAX_DEG 20) = 20 deg" in text
    wide = {**row, "ptch_lim_deg": 35.0, "theta_A_hi_deg": 30.0}
    assert "configured bound" in bound_reason(wide, "theta_A:hi")
    assert bound_reason(row, "spin_min:lo") == "" and bound_reason({}, "theta_A:hi") == ""


def test_the_c_chapter_carries_seed_parameters_the_ridge_sentence_and_bound_reason(tmp_path):
    c2_run(tmp_path)
    params = list(csv.DictReader((tmp_path / "C_params.csv").open(encoding="utf-8")))
    params[0].update({"at_bound": "theta_A:hi", "theta_A_cfg_hi_deg": "30", "ptch_lim_deg": "20",
                      "theta_A_hi_deg": "20"})
    write_table(tmp_path, "C_params", params)
    write_table(tmp_path, "C_search", seeds(
        [1.4796, 1.4789, 1.4795], T_bt_s=[1.2, 1.6, 2.1], theta_A_deg=[19.5, 19.6, 19.6],
        spin_min=[0.11, 0.11, 0.11], Q_TRANS_DECEL_ms2=[2.03, 2.03, 2.04]))
    c = section(build_html(load(tmp_path)), "c")
    assert "at a search bound: theta_A:hi (the upper bound of theta_A is the pitch limit" in c
    assert "min(30, PTCH_LIM_MAX_DEG 20) = 20 deg" in c
    assert "T_bt_s = 1.20 s" in c and "T_bt_s = 2.10 s" in c  # the parameters are in the tooltip
    assert "a flat ridge, not a single optimum" in c


def test_a_dot_chart_draws_the_reference_and_widens_the_axis_for_it():
    items = [("a", 10.0, "ta", "s2"), ("b", 11.0, "tb", "s2")]
    plain = ch.dot_chart(items, "x", "x", zero=False)
    refd = ch.dot_chart(items, "x", "x", zero=False, refs=[(25.0, "nominal a"), None])
    assert "mk ref hollow" not in plain and refd.count("mk ref hollow") == 1
    def top(svg):  # the largest tick label of the x axis
        return max(float(v) for v in re.findall(r'class="tick"[^>]*>([-0-9.]+)<', svg))

    assert "nominal a" in refd and top(plain) < 15 and top(refd) >= 25


# ---- F (the FT chapter) -----------------------------------------------------------------
def ft_hist_rows(case, wait_cap, fade_cap, scenario="S-07", v=17.0):
    rows = []
    for i in range(8):
        mode = "fwd_wait" if i < 3 else ("fwd_fade" if i < 6 else "fwd_done")
        cap = wait_cap if mode == "fwd_wait" else (fade_cap if mode == "fwd_fade" else 15.0)
        rows.append({"scenario": scenario, "case": case, "t": float(i), "mode": mode,
                     "V_eas": 3.0 * i, "h_rel": 0.1 * i, "theta": 2.0 * i, "theta_cap": cap,
                     "delta_r": 0.5 - 0.05 * i, "delta_p": 0.1 * i})
    return rows


def ft_run(tmp_path, with_runs=True):
    write_table(tmp_path, "FT_params", [
        {"case": "optimum", "airspeed_min_eas_ms": v, "q_transition_s": 17.0, "tkoff_thr_max": 1.0,
         "tran_pit_max_deg": 6.0, "J": 0.759, "n_flags": 0.0, "worst_stall_margin": 1.246}
        for v in (18.0,)] + [
        {"case": "register", "airspeed_min_eas_ms": 17.0, "q_transition_s": 5.0,
         "tkoff_thr_max": 0.8, "tran_pit_max_deg": 3.0, "J": 51.3, "n_flags": 5.0}])
    write_table(tmp_path, "FT_history", [
        *ft_hist_rows("optimum", 6.0, 14.0), *ft_hist_rows("register", 3.0, 8.0),
        *ft_hist_rows("ardupilot_default", 3.0, 8.0)])
    write_record(tmp_path)


def test_the_ft_history_draws_every_cases_own_cap_line(tmp_path):
    ft_run(tmp_path)
    ft = section(build_html(load(tmp_path)), "ft")
    for case in ("optimum", "register", "ArduPilot default"):
        assert f"pitch cap ({case})" in ft, case
    assert "pitch cap (optimum)" in ft and "pitch cap</" not in ft.replace("pitch cap (", "")


# ---- F2 -------------------------------------------------------------------------------------
def test_rule_labels_are_stacked_with_a_minimum_gap():
    rows = ch.stack_labels([100.0, 101.0, 102.5, 180.0], 11.0, 22.0, 200.0)
    assert rows[3] == 180.0  # a lone label stays on its rule
    near = sorted(rows[:3])
    assert all(b - a >= 11.0 - 1e-9 for a, b in zip(near, near[1:], strict=False))
    low = ch.stack_labels([196.0, 197.0, 198.0], 11.0, 22.0, 200.0)  # pushed back up at the bottom
    assert max(low) <= 200.0 and sorted(low)[1] - sorted(low)[0] >= 11.0 - 1e-9
    assert ch.stack_labels([], 11.0, 0.0, 10.0) == []


def test_the_airspeed_panel_starts_at_zero_with_labels_at_the_right_and_case_marks():
    t = [float(i) for i in range(10)]
    s = ch.TSeries("optimum", "s1", "optimum", t, {"V": [10.0 + i for i in range(10)]},
                   marks=[(3.0, "sq", "optimum: fade starts at 3.0 s"),
                          (6.0, "dia", "optimum: done at 6.0 s")])
    p = ch.Panel("V", "airspeed", "m/s", 2, hlines=[(17.0, "A"), (17.2, "B"), (17.4, "C")],
                 ymin=0.0, marks=True)
    svg = ch.panel_figure("f", [p], [s])
    assert "optimum: fade starts at 3.0 s" in svg and "optimum: done at 6.0 s" in svg
    ticks = [float(v) for v in re.findall(r'class="tick"[^>]*>([-0-9.]+)<', svg)]
    assert min(ticks) == 0.0
    ys = [float(m) for m in re.findall(r'<text class="dl" x="[0-9.]+" y="([0-9.]+)" '
                                       r'text-anchor="end">[ABC]<', svg)]
    assert len(ys) == 3 and all(b - a >= 11.0 - 0.2 for a, b in zip(sorted(ys), sorted(ys)[1:],
                                                                    strict=False))
    plain = ch.panel_figure("f", [ch.Panel("V", "airspeed", "m/s")], [s])
    assert "fade starts" not in plain  # marks only where the panel asks


def test_the_stall_rule_is_the_heaviest_members_not_the_mean():
    from sagetrans.io.view.chapters import heaviest_stall

    runs = [{"V_s_eas_ms": 12.0}, {"V_s_eas_ms": 14.0}, {"V_s_eas_ms": 13.0}, {"V_s_eas_ms": ""}]
    v, lab = heaviest_stall(runs, 1.2)
    assert abs(v - 16.8) < 1e-9 and lab == "1.2 x stall speed (heaviest member)"
    assert heaviest_stall([{}], 1.2) is None


AXES = {"altitude_m": [0, 5000], "isa_offset_k": [-10, 20], "mass_factor": [0.8, 1.15],
        "headwind_ms": [-5, 5], "soc": [0.2, 1.0]}


def test_the_scenario_caption_says_where_it_lies_against_the_ensemble():
    from sagetrans.io.view.chapters import ft_scenario_caption

    mid = {"altitude_m": 1500, "isa_offset_k": 0, "mass_factor": 1.0, "soc": 0.8,
           "headwind_ms": 0, "note": "reference altitude"}
    s = ft_scenario_caption("S-07", mid, AXES, [], [])
    assert "altitude 1500 m" in s and "ISA offset +0 K" in s and "(reference altitude)" in s
    assert "inside the ensemble's range on every axis." in s and "OUTSIDE" not in s
    corner = {**mid, "altitude_m": 5000, "mass_factor": 1.15, "isa_offset_k": 10, "soc": 0.5}
    s = ft_scenario_caption("S-03", corner, AXES, [], [])
    assert "inside the ensemble's range on every axis, at the edge of the range on altitude, " \
        "mass factor." in s
    out = ft_scenario_caption("X", {**mid, "isa_offset_k": 30, "headwind_ms": -8}, AXES, [], [])
    assert "OUTSIDE the ensemble's range on ISA offset (30 against -10 to 20), headwind (-8" in out
    hist = [{"scenario": "S-03", "case": "ardupilot_default", "h_rel": h}
            for h in (0.0, -13.2, -4.0)]
    runs = [{"case": "ardupilot_default", "loss_m": m} for m in (1.0, 3.9, 2.0)]
    s = ft_scenario_caption("S-03", corner, AXES, hist, runs)
    assert "default loses 13.2 m of altitude here" in s and "at most 3.9 m" in s
    assert ft_scenario_caption("S-09", None, AXES, [], []) == "Scenario S-09."


def test_the_ft_chapter_has_the_airspeed_rules_marks_and_scenario_caption(tmp_path):
    ft_run(tmp_path)
    write_table(tmp_path, "FT_runs", [
        {"case": "optimum", "set": "ensemble", "member": float(i), "V_s_eas_ms": 12.0 + i,
         "loss_m": 1.0, "flags": ""} for i in range(3)])
    write_record(tmp_path, context={"ft_scenarios": {
        "scenarios": {"S-07": {"altitude_m": 1500, "isa_offset_k": 0, "mass_factor": 1.0,
                               "soc": 0.8, "headwind_ms": 0, "note": "reference"}},
        "axes": AXES}})
    ft = section(build_html(load(tmp_path)), "ft")
    assert "1.2 x stall speed (heaviest member)" in ft and "stall speed</" not in ft
    assert "AIRSPEED_MIN optimum (18)" in ft and "AIRSPEED_MIN register (17)" in ft
    assert "Scenario S-07: altitude 1500 m" in ft and "inside the ensemble" in ft
    assert "fade starts at 3.0 s" in ft and "done at 6.0 s" in ft and "fade (optimum)" in ft


# ---- F3 -------------------------------------------------------------------------------------
BEST = {"airspeed_min": 18.0, "q_transition": 17.0, "tkoff_thr_max": 1.0, "tran_pit_max": 6.0}


def fake_eval(flagged):
    """summarise_ft-like rows: `flagged(airspeed_min, scale)` -> number of flagged members."""
    def ev(prm, scale):
        n = flagged(prm["airspeed_min"], scale)
        return {"J": 0.76 + 0.01 * (prm["airspeed_min"] - 18.0) + 20.0 * n,
                "worst_stall_margin": prm["airspeed_min"] / 14.5 * scale, "n_flags": float(n),
                "n_stall_margin": float(n), "n_stall": 0.0, "n_alt_loss": 0.0}
    return ev


def test_ft_robust_has_a_row_per_airspeed_and_scale_and_names_the_flags():
    rows = stages.robust_rows(BEST, [-1, 0, 1, 2], [0.9, 1.0, 1.1], fake_eval(lambda v, s: 0))
    assert len(rows) == 12
    assert [r["airspeed_min_eas_ms"] for r in rows[::3]] == [17.0, 18.0, 19.0, 20.0]
    assert {r["cl_max_scale"] for r in rows} == {0.9, 1.0, 1.1} and rows[0]["flags"] == ""
    bad = stages.robust_rows(BEST, [0], [0.9], fake_eval(lambda v, s: 2))
    assert bad[0]["flags"] == "stall_margin" and bad[0]["n_flags"] == 2.0


def test_optimum_robust_is_the_smallest_airspeed_clean_at_every_scale():
    # 17 is flagged everywhere, 18 only at C_L,max x 0.9, 19 and 20 are clean
    rows = stages.robust_rows(BEST, [-1, 0, 1, 2], [0.9, 1.0, 1.1],
                              fake_eval(lambda v, s: int(v < 17.5 or (v < 18.5 and s < 0.95))))
    r = stages.optimum_robust_row(rows, BEST)
    assert r["case"] == "optimum_robust" and r["robust_found"] is True
    assert r["airspeed_min_eas_ms"] == 19.0 and r["q_transition_s"] == 17.0
    assert r["tkoff_thr_max"] == 1.0 and r["tran_pit_max_deg"] == 6.0 and r["n_flags"] == 0.0
    j_opt = 0.76 + 20.0 * 0  # the optimum at the nominal scale (1.0) is clean
    assert abs(r["J_optimum_same_ensemble"] - j_opt) < 1e-9
    assert abs(r["delta_J"] - 0.01) < 1e-9 and abs(r["J"] - 0.77) < 1e-9
    assert abs(r["worst_stall_margin"] - 19.0 / 14.5 * 0.9) < 1e-9  # the worst over the scales


def test_with_no_flag_free_airspeed_the_row_says_so_and_has_no_settings():
    rows = stages.robust_rows(BEST, [-1, 0, 1, 2], [0.9, 1.0], fake_eval(lambda v, s: 1))
    r = stages.optimum_robust_row(rows, BEST)
    assert r["robust_found"] is False and "airspeed_min_eas_ms" not in r and "delta_J" not in r
    assert r["case"] == "optimum_robust" and r["cl_max_scales"] == "0.9,1"


def test_when_the_optimum_is_already_robust_the_row_equals_it():
    rows = stages.robust_rows(BEST, [-1, 0, 1, 2], [0.9, 1.0, 1.1],
                              fake_eval(lambda v, s: int(v < 17.5)))
    r = stages.optimum_robust_row(rows, BEST)
    assert r["airspeed_min_eas_ms"] == 18.0 and r["delta_J"] == 0.0
    assert r["J"] == r["J_optimum_same_ensemble"]


def test_the_robust_settings_are_refused_out_of_range():
    import pytest
    from pydantic import ValidationError

    from sagetrans.io.settings import FTRobustCfg

    FTRobustCfg()  # the defaults are valid
    for bad in ({"airspeed_min_steps": []}, {"airspeed_min_steps": [-1, 11]},
                {"airspeed_min_steps": [float("nan")]}, {"cl_max_scale": []},
                {"cl_max_scale": [0.4]}, {"cl_max_scale": [1.6]}, {"cl_max_scale": [0.9, 0.0]},
                {"unknown": 1}):
        with pytest.raises(ValidationError):
            FTRobustCfg(**bad)


def robust_table(flag_at):
    rows = []
    for v in (17.0, 18.0, 19.0, 20.0):
        for s in (0.9, 1.0, 1.1):
            n = flag_at(v, s)
            rows.append({"airspeed_min_eas_ms": v, "cl_max_scale": s, "J": 0.76 + 20.0 * n,
                         "worst_stall_margin": 1.25 if not n else 1.18, "n_flags": float(n),
                         "flags": "stall_margin" if n else ""})
    return rows


def test_the_robust_sentence_is_computed_from_the_table():
    from sagetrans.io.view.chapters import robust_statement

    opt = {"airspeed_min_eas_ms": 18.0}
    rob = {"case": "optimum_robust", "robust_found": True, "airspeed_min_eas_ms": 19.0,
           "J": 0.77, "delta_J": 0.01, "delta_J_in_seed_spreads": 1.3}
    rows = robust_table(lambda v, s: int(v < 17.5 or (v < 18.5 and s < 0.95)))
    s = robust_statement(rows, opt, rob)
    assert s.startswith("AIRSPEED_MIN 19 m/s is flag-free at all 3 C_L,max scales (0.9, 1, 1.1); ")
    assert "the optimum's 18 m/s is flag-free only at 1, 1.1." in s
    assert "+0.010 against the optimum on the same ensemble (+1.3 seed spreads)" in s
    assert "reported, not recommended" in s
    none = robust_statement(robust_table(lambda v, s: 1), opt, {"robust_found": False})
    assert none.startswith("No AIRSPEED_MIN in the table is flag-free at every C_L,max scale")
    assert "flag-free at none of them" in none
    same = robust_statement(robust_table(lambda v, s: int(v < 17.5)), opt,
                            {"robust_found": True, "airspeed_min_eas_ms": 18.0})
    assert "AIRSPEED_MIN 18 m/s, the optimum, is flag-free at all 3" in same
    assert robust_statement([], opt, None) == ""


def test_the_ft_chapter_has_the_robust_table_the_tile_and_the_sentence(tmp_path):
    ft_run(tmp_path)
    flagged = robust_table(lambda v, s: int(v < 17.5 or (v < 18.5 and s < 0.95)))
    write_table(tmp_path, "FT_robust", flagged)
    rows = list(csv.DictReader((tmp_path / "FT_params.csv").open()))
    rows.append({"case": "optimum_robust", "robust_found": "True", "airspeed_min_eas_ms": "19",
                 "J": "0.77", "delta_J": "0.01", "delta_J_in_seed_spreads": "1.3"})
    write_table(tmp_path, "FT_params", rows)
    ft = section(build_html(load(tmp_path)), "ft")
    assert "Stall margin spare" in ft and "1.246 vs 1.2 (3.8%)" in ft
    assert "Stall margin robustness" in ft and "flag-free only at 1, 1.1." in ft


# ---- F4 -------------------------------------------------------------------------------------
def tick_labels(svg):
    return re.findall(r'class="tick"[^>]*>([-0-9.e+]+)<', svg)


def test_a_log_axis_separates_neighbouring_costs_and_has_no_zero():
    s = ch.Series("J", "s1", [(17.0, 20.8), (18.0, 0.7588), (19.0, 0.7604), (20.0, 51.0)])
    lin = ch.xy_chart([s], "v", "J", "lin")
    log = ch.xy_chart([s], "v", "J (log)", "log", ylog=True)

    def cy(svg):
        return [float(v) for v in re.findall(r'<circle class="mk s1"[^>]* cy="([0-9.]+)"', svg)]

    gap_lin = abs(cy(lin)[1] - cy(lin)[2])
    gap_log = abs(cy(log)[1] - cy(log)[2])
    assert gap_log > 2 * gap_lin  # 18 and 19 can be told apart on the log axis
    labels = tick_labels(log)
    assert "0" not in labels and "0.1" in labels and "1" in labels and "100" in labels
    assert "0.2" in labels and "5" in labels  # decades subdivided at 2 and 5


def test_a_symmetric_log_axis_keeps_small_rows_visible_and_labels_real_units():
    items = [("small a", 0.002, "", "s1"), ("small b", -0.003, "", "s1"), ("big", 20.0, "", "s2")]
    lin = ch.dot_chart(items, "dJ", "lin")
    sym = ch.dot_chart(items, "dJ", "sym", sym=0.003)

    def cx(svg):
        return [float(v) for v in re.findall(r'<circle class="mk s[12]"[^>]* cx="([0-9.]+)"', svg)]

    assert abs(cx(lin)[0] - cx(lin)[1]) < 3  # flattened
    assert abs(cx(sym)[0] - cx(sym)[1]) > 15  # separate
    labels = tick_labels(sym)
    assert "0" in labels and "30" in labels  # zero is always labelled, the rest in data units
    assert all(float(v) in (0.0, 0.003, 0.03, 0.3, 3.0, 30.0, -0.003, -0.03, -0.3, -3.0, -30.0)
               for v in labels)


def test_the_swarm_keeps_overlapping_marks_apart():
    xs = [100.0, 100.5, 101.0, 101.2, 150.0]
    ys = ch.swarm(xs)
    assert ys[0] == 0.0 and ys[-1] == 0.0  # a lone mark stays on the line
    pts = list(zip(xs, ys, strict=True))
    for i, (x, y) in enumerate(pts):
        for x2, y2 in pts[:i]:
            assert math.hypot(x - x2, y - y2) >= 10.0 - 1e-9


def ft_sens_rows():
    base = {"parameter": "base", "value": "", "J": 0.7588, "n_flags": 0.0}
    rows = [base]
    for p, v, j, n in (("rotor_lag_s", 0.5, 0.7602, 0), ("z_acc_p", 0.015, 0.7610, 0),
                       ("TRIM_THROTTLE", 0.35, 0.7590, 0), ("CL_max_scale", 0.9, 20.8, 2)):
        rows.append({"parameter": p, "value": v, "J": j, "n_flags": float(n)})
    return rows


def test_the_sensitivity_sentence_is_computed_and_can_say_no():
    from sagetrans.io.view.chapters import sensitivity_statement

    s = sensitivity_statement(ft_sens_rows(), 0.7588, 0.0039)
    assert "seed spread is 0.39% of J, 0.0030 in J" in s
    assert "Apart from C_L,max, no input moves J by more than that" in s
    assert "the largest is z_acc_p = 0.015 at +0.0022" in s
    tight = sensitivity_statement(ft_sens_rows(), 0.7588, 0.0001)  # a tiny spread: not true now
    assert "DO move J by more than that" in tight and "rotor_lag_s = 0.5" in tight
    assert "CL_max_scale" not in tight
    assert sensitivity_statement(ft_sens_rows(), 0.7588, float("nan")) == ""


def grid_rows_ft():
    def g(q, j, **kw):
        return {"q_transition_s": q, "J_snapped": j, "n_flags": 0.0, **kw}

    return [g(2.0, 0.90, n_flags=3.0), g(5.0, 0.7700), g(10.0, 0.7620), g(17.0, 0.7588),
            g(20.0, 0.7600)]


def test_the_q_transition_sentence_gives_range_register_plateau_and_the_second_order_claim():
    from sagetrans.io.view.chapters import grid_statement

    params = [{"case": "optimum", "J": 0.7588, "J_start": 51.3},
              {"case": "register", "q_transition_s": 5.0}]
    s = grid_statement(grid_rows_ft(), params, 0.0039)
    assert "J runs from 0.759 (Q_TRANSITION_MS 17 s) to 0.900 (2 s), 18.6% above the lowest" in s
    assert "At the register's Q_TRANSITION_MS (5 s) J is 0.770, +1.5% against the winner's." in s
    assert "2 of 5 points lie within the seed spread (0.39%) of the best: " \
           "Q_TRANSITION_MS 17, 20 s." in s
    assert "so with the other three re-tuned the fade length is a second-order choice." in s
    wide = grid_statement(grid_rows_ft(), [{"case": "optimum", "J": 0.7588, "J_start": 0.95}],
                          0.0039)
    assert "not a second-order choice" in wide
    off = grid_statement(grid_rows_ft(), [{"case": "register", "q_transition_s": 7.0}], 0.0039)
    assert "(7 s) is not on the grid" in off
    assert grid_statement(grid_rows_ft()[:1], params, 0.004) == ""


def test_the_spread_sentence_counts_default_members_and_compares_energy():
    from sagetrans.io.view.chapters import share_statement

    dmem = [{"stall_margin_min": m, "energy_kJ": 30.0} for m in (0.8, 0.95, 1.1, 1.3)]
    opt = [{"stall_margin_min": 1.25, "energy_kJ": 33.0}] * 3
    reg = [{"stall_margin_min": 1.0, "energy_kJ": 30.0}] * 3
    s = share_statement(dmem, opt, reg)
    assert "2 of 4 ArduPilot-default members (50%) fly below a stall margin of 1.0 and 3 (75%) " \
           "below 1.2." in s
    assert "optimum's median energy is 33.00 kJ: +10.0% against the default's 30.00 kJ and " \
           "+10.0% against the register's 30.00 kJ" in s
    assert share_statement([], [], []) == ""


def test_the_tkoff_thr_max_chip_says_why_it_is_on_the_bound():
    from sagetrans.io.view.chapters import ft_bound_reason

    row = {"tkoff_thr_max_hi": 1.0, "thr_max": 1.0}
    s = ft_bound_reason(row, "tkoff_thr_max:hi")
    assert "upper bound of TKOFF_THR_MAX is 100%" in s and "THR_MAX in the register is 100%" in s
    assert ft_bound_reason(row, "airspeed_min:hi") == ""
    assert ft_bound_reason({}, "tkoff_thr_max:hi") == ""


def ft_f4_run(tmp_path):
    ft_run(tmp_path)
    write_table(tmp_path, "FT_runs", [
        {"case": cs, "set": "core", "member": float(i), "V_s_eas_ms": 12.0 + i,
         "loss_m": 1.0 + i, "gain_m": 0.5, "stall_margin_min": 1.0 + 0.1 * i + k,
         "energy_kJ": 30.0 + k * 9, "vh_end_ms": 0.5 - 0.1 * i, "h_proj_rel_m": -2.0 - i,
         "proj_loss_m": 0.0, "wing_share_handover": 0.4 + 0.01 * i + k * 0.01,
         "V_handover_eas_ms": 18.0, "flags": ""}
        for k, cs in enumerate(("optimum", "register", "ardupilot_default")) for i in range(6)])
    write_table(tmp_path, "FT_search", [
        {"seed": float(s), "best_cost": 0.76 + 0.001 * s, "evaluations": 100.0,
         "airspeed_min_eas_ms": 18.0, "q_transition_s": 17.0, "tkoff_thr_max": 1.0,
         "tran_pit_max_deg": 6.0} for s in (1, 2, 3)])
    write_table(tmp_path, "FT_scan", [
        {"airspeed_min_eas_ms": v, "J": j, "n_flags": float(n), "flags": f,
         "q_transition_s": 17.0, "tkoff_thr_max": 1.0, "tran_pit_max_deg": 6.0}
        for v, j, n, f in ((16.0, 5.0, 3, "stall_margin"), (17.0, 1.4, 2, "stall_margin,stall"),
                           (18.0, 0.7588, 0, ""), (19.0, 0.7604, 0, ""),
                           (20.0, 3.0, 1, "alt_loss"))])
    write_table(tmp_path, "FT_sens", ft_sens_rows())
    write_record(tmp_path, search_spreads={"FT": 0.0039},
                 context={"mission": {"fwd_alt_loss_max_m": 10.0, "fwd_alt_gain_max_m": 10.0}})


def test_the_ft_chapter_has_the_f4_figures(tmp_path):
    ft_f4_run(tmp_path)
    ft = section(build_html(load(tmp_path)), "ft")
    # 5: the end of the fade
    assert "sink rate [m/s] (negative = climbing)" in ft
    assert "The loss limit (-10 m) is outside the range shown" in ft  # zoomed to the data
    # 6: one strip per case, AIRSPEED_MIN in the label, every member a dot
    assert "optimum (AIRSPEED_MIN 18)" in ft and "register (AIRSPEED_MIN 17)" in ft
    # 7: final cost by seed, with the four settings in the tooltip
    assert "Final cost by seed" in ft and "Search convergence" not in ft
    assert "seed = 2; best_cost = 0.762" in ft and "tran_pit_max_deg = 6" in ft
    # 1: the scan on a log axis, the hollow tooltip names the flags
    assert "cost J (log scale)" in ft and "stall_margin,stall" in ft
    assert "Flag-free: 18, 19 m/s." in ft
    # 3: the sensitivity in absolute change with the sentence
    assert "Absolute change of the cost J" in ft and "Apart from C_L,max, no input moves J" in ft
    assert "[2 failing]" in ft
    # 4: the register's members, the default's members and the limits
    assert "member of the register start" in ft and "member of the ArduPilot default" in ft
    assert ">limit<" in ft and "limit 1.2" in ft
    assert "ArduPilot-default members" in ft and "median energy is 30.00 kJ" in ft


def test_direct_labels_move_off_each_other():
    placed = []
    first = ch.place_label(100.0, 100.0, "q = 2", (-9.0, 17.0), placed)
    second = ch.place_label(104.0, 101.0, "q = 2.25", (-9.0, 17.0), placed)
    third = ch.place_label(108.0, 102.0, "q = 2.5", (-9.0, 17.0), placed)
    assert first == (-9.0, 17.0) and len({first, second, third}) == 3
    far = ch.place_label(400.0, 300.0, "alone", (8.0, -8.0), placed)
    assert far == (8.0, -8.0)  # a label with room keeps its preferred place


# ---- found by looking at the figures of the tiny run ----------------------------------------
def test_a_wide_panel_spans_both_columns_and_a_long_strip_label_has_room():
    s = ch.TSeries("a", "s1", "a", [0.0, 1.0], {"V": [1.0, 2.0], "h": [0.0, 1.0]})
    svg = ch.panel_figure("f", [ch.Panel("V", "v", wide=True), ch.Panel("h", "h")], [s])
    assert svg.count('class="pan wide"') == 1 and svg.count('class="pan"') == 1
    assert 'viewBox="0 0 720 200"' in svg and 'viewBox="0 0 350 200"' in svg
    long = "ArduPilot default (AIRSPEED_MIN 9) and more words"
    strip = ch.strip_chart([(long, [(1.0, "", False, "s1"), (2.0, "", False, "s1")], None)], "x")
    pl = float(re.search(r'class="axlab" x="([0-9.]+)"', strip).group(1))
    assert pl >= 6.0 * len(long) - 12  # the label ends at x = pl - 12 and starts inside the figure


def test_a_strip_row_is_as_tall_as_its_swarm_and_its_rule_label_stays_inside():
    few = ch.strip_chart([("m", [(1.0, "", False, "s1"), (9.0, "", False, "s1")], None)], "x")
    dots = [(1.0, "", False, "s1")] * 20 + [(9.0, "", False, "s1")]
    many = ch.strip_chart([("m", dots, None)], "x")

    def height(svg):
        return float(re.search(r'viewBox="0 0 720 ([0-9.]+)"', svg).group(1))

    assert height(many) > height(few)
    right = ch.strip_chart([("m", [(1.0, "", False, "s1")], None)], "x",
                           extras=[([], (9.0, "tolerance 1.76"))])
    left = ch.strip_chart([("m", [(9.0, "", False, "s1")], None)], "x",
                          extras=[([], (1.0, "tolerance 1.76"))])
    assert 'text-anchor="end">tolerance 1.76<' in right  # near the right edge: ends at the rule
    assert 'text-anchor="start">tolerance 1.76<' in left


# ---- every new table in one small folder ---------------------------------------------------------
def test_the_viewer_and_the_workbook_build_from_a_small_folder_with_every_new_table(tmp_path):
    from openpyxl import load_workbook

    from sagetrans.io.view.workbook import write_workbook

    synthetic_run(tmp_path)
    exposure = {"stall_members": 4.0, "stall_alpha_peak_deg_max": 11.7,
                "stall_members_above_peak": 0.0, "stall_t_above_peak_s_max": 0.0,
                "stall_t_above_limit_s_max": 0.4, "stall_wing_share_min": 0.62}
    held = {"alpha_enforced": True, "alpha_limit_deg": 10.99, "p95_stop_time_s": 18.0,
            "alpha_s_deg": 15.7, "blend_M": 70.0, "cl_max": 1.444, "alpha_peak_deg": 13.43,
            "rotor_constant_coeffs": True, "max_overshoot_m": 1.4, "max_end_sink_ms": 1.2}
    write_table(tmp_path, "C_params", [
        {"case": "optimum", "T_bt_s": 2.0, "theta_A_deg": 9.0, "spin_min": 0.1,
         "Q_TRANS_DECEL_ms2": 1.05, "J": 2.0, "p95_distance_m": 237.0, "mean_energy_kJ": 36.6,
         **held, **exposure},
        {"case": "ardupilot_default", "T_bt_s": 3.0, "theta_A_deg": 30.0, "spin_min": 0.15,
         "Q_TRANS_DECEL_ms2": 2.0, "J": 61.5, "p95_distance_m": 121.0, "mean_energy_kJ": 20.3}])
    write_table(tmp_path, "C_ref_params", [
        {"case": "optimum", "T_bt_s": 1.6, "theta_A_deg": 19.6, "spin_min": 0.11,
         "Q_TRANS_DECEL_ms2": 1.79, "J": 1.5, "p95_distance_m": 140.0, "mean_energy_kJ": 21.0,
         "p95_stop_time_s": 10.4, "alpha_enforced": False, **{**exposure,
                                                              "stall_members_above_peak": 4.0}}])
    write_table(tmp_path, "C_ref_runs", [{"member": 0, "energy_kJ": 21.0}])
    write_table(tmp_path, "C_ref_history", hist_rows("C optimum", "core", "S-01", -1))
    write_table(tmp_path, "C_alpha_curve", [
        {"kind": "unconstrained", "alpha_frac": "", "limit_deg": "", "Q_TRANS_DECEL_ms2": 1.79,
         "p95_distance_m": 140.0, "p95_stop_time_s": 10.4, "mean_energy_kJ": 21.0},
        {"kind": "enforced", "alpha_frac": 1.0, "limit_deg": 15.7, "Q_TRANS_DECEL_ms2": 1.3,
         "p95_distance_m": 200.0, "p95_stop_time_s": 14.0, "mean_energy_kJ": 30.0},
        {"kind": "enforced", "alpha_frac": 0.7, "limit_deg": 10.99, "Q_TRANS_DECEL_ms2": 1.05,
         "p95_distance_m": 237.0, "p95_stop_time_s": 18.0, "mean_energy_kJ": 36.6}])
    html = build_html(load(tmp_path))
    c = section(html, "c")
    for text in ("The enforced configuration and the unconstrained reference",
                 "the model ignores the stall; reference only; never a recommendation",
                 "What this result rests on"):
        assert text in c, text
    assert 'id="b1"' not in html and "B1" not in c
    out = tmp_path / "tables.xlsx"
    write_workbook(load(tmp_path), out)
    sheets = " ".join(load_workbook(out).sheetnames)
    assert "C_alpha_curve" in sheets and "C_ref_params" in sheets
